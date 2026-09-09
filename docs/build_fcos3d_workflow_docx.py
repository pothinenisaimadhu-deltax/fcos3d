from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'fcos3d_workflow_interactive.html'
OUTPUT = ROOT / 'fcos3d_pkl_data_train_test_explained.docx'


def set_font(run, name='Calibri', size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.bold = bold


def set_style(style, name, size, color, before, after, line):
    style.font.name = name
    style._element.rPr.rFonts.set(qn('w:ascii'), name)
    style._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(*color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def label_para(doc, label, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(5)
    lead = p.add_run(f'{label}: ')
    set_font(lead, size=10.5, color=(31, 77, 120), bold=True)
    body = p.add_run(text)
    set_font(body, size=10.5, color=(36, 55, 72))
    return p


def parse_steps():
    source = SOURCE.read_text(encoding='utf-8')
    pattern = re.compile(
        r"\['([^']*)',(\d),'([^']*)','([^']*)','([^']*)','([^']*)','([^']*)'\]"
    )
    rows = []
    for match in pattern.finditer(source):
        step_id, lane, title, file_path, flow, input_text, output_text = match.groups()
        rows.append(dict(
            id=step_id, lane=int(lane), title=title, file=file_path,
            flow=flow, input=input_text, output=output_text))
    if len(rows) != 30:
        raise RuntimeError(f'Expected 30 workflow steps; found {len(rows)}.')
    return rows


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    set_font(run, size=9, color=(98, 120, 135))
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def main():
    rows = parse_steps()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles['Normal']
    set_style(normal, 'Calibri', 11, (36, 55, 72), 0, 6, 1.25)
    for style_name, size, color, before, after in [
        ('Heading 1', 16, (46, 116, 181), 18, 10),
        ('Heading 2', 13, (46, 116, 181), 14, 7),
        ('Heading 3', 12, (31, 77, 120), 10, 5),
    ]:
        set_style(doc.styles[style_name], 'Calibri', size, color, before, after, 1.25)

    subtitle = doc.styles.add_style('SubtitleCustom', WD_STYLE_TYPE.PARAGRAPH)
    set_style(subtitle, 'Calibri', 11, (98, 120, 135), 0, 16, 1.25)

    header = section.header.paragraphs[0]
    header.text = 'FCOS3D PKL Workflow Reference'
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, size=9, color=(98, 120, 135), bold=True)
    page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run('FCOS3D: Data Preparation, Training, and Testing')
    set_font(title_run, size=24, color=(11, 37, 69), bold=True)
    sub = doc.add_paragraph('Camera-only NuScenes workflow using the active PKL manifest route', style='SubtitleCustom')
    doc.add_paragraph('Purpose: a function-level reference for tracing how raw NuScenes-format data becomes an FCOS3D checkpoint, then 3D predictions and evaluation output.')

    doc.add_heading('How to read this document', level=1)
    doc.add_paragraph('The workflow contains 30 ordered steps. Each entry gives the file that owns the behavior, the main function sequence, the values or artifacts entering that stage, and the result passed to the next stage. The document intentionally excludes the removed JSON preparation alternatives.')
    label_para(doc, 'Data path', 'Raw NuScenes tables and camera images -> v2 train/validation PKL manifests -> six camera-image records per keyframe -> packed training batches.')
    label_para(doc, 'Training path', 'FCOS3D configuration -> model initialization -> image features -> FCOS3D dense predictions -> losses -> optimizer updates -> checkpoints.')
    label_para(doc, 'Testing path', 'Checkpoint -> test runner -> camera-space 3D boxes -> score/NMS filtering -> NuScenes-format JSON -> metrics or exported custom-data predictions.')

    overview = [
        ('Data preparation', 'D01-D10', 'Builds and consumes v2 PKL manifests; expands every keyframe into camera-specific examples.'),
        ('Training and validation', 'T01-T10', 'Loads FCOS3D, derives dense 3D targets, optimizes the model, and saves checkpoints.'),
        ('Testing and outputs', 'E01-E10', 'Loads a checkpoint, decodes 3D detections, evaluates or exports them, and optionally renders them.'),
    ]
    doc.add_heading('Workflow map', level=1)
    for name, identifier, explanation in overview:
        p = doc.add_paragraph(style='Heading 2')
        p.add_run(f'{identifier}  {name}')
        doc.add_paragraph(explanation)

    names = ['Data preparation', 'Training and validation', 'Testing and outputs']
    lead = [
        'This section follows the checked-in PKL preparation route. The result is a v2 NuScenes-compatible train/validation manifest with per-camera annotations.',
        'This section shows how camera-image samples become FCOS3D losses, parameter updates, validation predictions, and checkpoint files.',
        'This section shows how a saved checkpoint produces camera-space boxes, NuScenes result JSON, metrics, and optional visual overlays.',
    ]
    for lane in range(3):
        doc.add_page_break()
        doc.add_heading(names[lane], level=1)
        doc.add_paragraph(lead[lane])
        for row in [item for item in rows if item['lane'] == lane]:
            doc.add_heading(f"{row['id']}  {row['title']}", level=2)
            label_para(doc, 'Source file', row['file'])
            label_para(doc, 'Function flow', row['flow'])
            label_para(doc, 'Input', row['input'])
            label_para(doc, 'Output and effect', row['output'])

    doc.add_page_break()
    doc.add_heading('Important operational checks', level=1)
    checks = [
        'The configured annotation file and the evaluator annotation file must identify the same split and use compatible v2 data_list records.',
        'mv_image_based makes six camera-image items from a keyframe. A batch size therefore counts images rather than whole six-camera scenes.',
        'Resize3D must keep the image, centers_2d, and cam2img in the same geometric frame; changing it can invalidate training/inference geometry.',
        'Official NuScenes metrics require matching official tokens and ground truth. Custom tokens should use format_only=True with jsonfile_prefix for result export.',
        'The test entry guard checks FCOSMono3D and NuScenesDataset before cfg-options are merged; testing a different dataset class requires intentional integration work.',
    ]
    for check in checks:
        doc.add_paragraph(check, style='List Bullet')

    doc.add_paragraph('Source snapshot: C:\\Users\\saima\\d-code\\mmdetection3d, 08 Sep 2026.', style='SubtitleCustom')
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
